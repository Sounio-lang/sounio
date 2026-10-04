#!/usr/bin/env python3
"""tools/ledger/build_ledger.py -- the Sounio work, as one spreadsheet.

Marshalling only: this script reads files and the GitHub API and writes
cells. It computes no scientific number. Every number in the "Resultados"
sheet comes from docs/ledger/results.tsv and, with --verify, is re-checked by
running the command that produced it and looking for the printed value.

Sheets, and where each comes from:
  Resumo              formulas over the sheets below
  Linhas de pesquisa  docs/ledger/research_lines.tsv  (+ path existence check)
  Issues e PRs        GitHub REST, repos/Sounio-lang/sounio/issues (open)
  Limitacoes          docs/compiler/KNOWN_LIMITATIONS.md (KL-xx headings),
                      LANGUAGE_GAPS.md of sounio-uhs-coupled (G-xx headings),
                      docs/ledger/defects.tsv
  Resultados          docs/ledger/results.tsv (+ --verify)

Usage:
  python3 tools/ledger/build_ledger.py --out sounio_ledger.xlsx \
      [--uhs-repo ../sounio-uhs-coupled] [--gri-repo ../sounio-gri30-crossvalidation] \
      [--verify] [--no-github]

Fails closed: a results row whose producer file is missing is reported as
"Ausente", never as valid; a check whose producer exits non-zero, or does not
print the expected string, is "Falhou", never silently dropped.
"""
import argparse
import csv
import datetime as dt
import json
import os
import re
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
FONT = "Arial"
HDR = Font(name=FONT, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="111821")
BODY = Font(name=FONT, size=10)
INPUT = PatternFill("solid", fgColor="FFF2CC")  # cells a person fills in
LINE = Side(style="thin", color="D9D9D9")
GREEN = PatternFill("solid", fgColor="D9EAD3")
RED = PatternFill("solid", fgColor="F4CCCC")
GREY = PatternFill("solid", fgColor="EEEEEE")


def read_tsv(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def sheet(wb, title, headers, rows, widths, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = HDR, HFILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.font = BODY
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(bottom=LINE)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    return ws


def repo_path(name, args):
    return {
        "sounio": ROOT,
        "sounio-uhs-coupled": args.uhs_repo,
        "sounio-gri30-crossvalidation": args.gri_repo,
    }.get(name)


# ---------------------------------------------------------------- research
def research_rows(args):
    rows = []
    for r in read_tsv(os.path.join(ROOT, "docs/ledger/research_lines.tsv")):
        base = repo_path(r["repositorio"], args)
        if base is None:
            exists = "repo não informado"
        else:
            exists = "sim" if os.path.exists(os.path.join(base, r["caminho"])) else "NÃO"
        rows.append([r["id"], r["linha"], r["repositorio"], r["caminho"], exists,
                     r["artefato_principal"], r["estado_medido"], r["medido_em"],
                     r["proximo_passo"], r["prazo"], r["paper_ou_destino"]])
    return rows


# ---------------------------------------------------------------- github
def gh_open_items():
    items, page = [], 1
    while True:
        out = subprocess.run(
            ["gh", "api", f"repos/Sounio-lang/sounio/issues?state=open&per_page=100&page={page}"],
            capture_output=True, text=True, timeout=120)
        if out.returncode != 0:
            raise RuntimeError(out.stderr.strip()[:300])
        batch = json.loads(out.stdout)
        if not batch:
            return items
        items.extend(batch)
        page += 1


AREAS = ("compiler", "stdlib", "ci", "documentation")


def issue_rows():
    rows = []
    for it in gh_open_items():
        labels = [l["name"] for l in it.get("labels", [])]
        area = next((a for a in AREAS if a in labels), "")
        created = dt.datetime.strptime(it["created_at"][:10], "%Y-%m-%d").date()
        rows.append([it["number"], "PR" if it.get("pull_request") else "Issue", it["title"],
                     area, ", ".join(labels), it["user"]["login"], created, None,
                     "", it["html_url"]])
    rows.sort(key=lambda r: r[0], reverse=True)
    return rows


# ---------------------------------------------------------------- limitations
def kl_rows():
    rows = []
    path = os.path.join(ROOT, "docs/compiler/KNOWN_LIMITATIONS.md")
    for line in open(path, encoding="utf-8"):
        m = re.match(r"^### (KL-[0-9a-z]+) — (.*)$", line.strip())
        if m:
            title = m.group(2)
            rows.append([m.group(1), "KNOWN_LIMITATIONS.md", title, "",
                         "fechado" if "CLOSED" in title else "aberto",
                         "docs/compiler/KNOWN_LIMITATIONS.md", ""])
    return rows


def gap_rows(args):
    if not args.uhs_repo:
        return []
    path = os.path.join(args.uhs_repo, "LANGUAGE_GAPS.md")
    if not os.path.exists(path):
        return []
    rows = []
    for line in open(path, encoding="utf-8"):
        m = re.match(r"^## (G[0-9]+) — (.*)$", line.strip())
        if m:
            title = m.group(2)
            rows.append([m.group(1), "uhs-coupled LANGUAGE_GAPS.md", title, "",
                         "fechado" if "CLOSED" in title else "aberto",
                         "sounio-uhs-coupled/LANGUAGE_GAPS.md", ""])
    return rows


def defect_rows():
    return [[d["id"], d["area"], d["descricao"], d["motor"], d["estado"], d["rastreio"],
             d["reproducao"]] for d in read_tsv(os.path.join(ROOT, "docs/ledger/defects.tsv"))]


# ---------------------------------------------------------------- results
def run_producer(cmd, cwd):
    """Run one producer command; return (stdout+stderr, seconds, exit code)."""
    t0 = dt.datetime.now()
    p = subprocess.run(["bash", "-c", cmd], cwd=cwd,
                       capture_output=True, text=True, timeout=900,
                       env={**os.environ, "SOUNIO_STDLIB_PATH": os.path.join(ROOT, "stdlib")})
    return p.stdout + p.stderr, (dt.datetime.now() - t0).total_seconds(), p.returncode


def judge(espera, out, secs, rc):
    """Verificado only when the producer exited 0 AND printed the expected string.
    A producer that printed the value and then failed is Falhou: the string alone
    does not prove a successful run."""
    found = espera in out
    if rc == 0 and found:
        return "Verificado", f"{secs:.0f} s"
    why = []
    if rc != 0:
        why.append(f"saiu com código {rc}")
    if not found:
        why.append(f"não encontrou: {espera}")
    return "Falhou", f"{secs:.0f} s; " + "; ".join(why)


def results_rows(args, path=None):
    cache = {}
    rows = []
    for r in read_tsv(path or os.path.join(ROOT, "docs/ledger/results.tsv")):
        base = repo_path(r["repositorio"], args)
        verdict, note = "não verificado", ""
        if args.verify:
            if not r["comando"]:
                verdict, note = "sem produtor local", "registrado na fonte; não reexecutável aqui"
            elif base is None or not os.path.isdir(base):
                verdict, note = "repo ausente", f"passe --{'uhs' if 'uhs' in r['repositorio'] else 'gri'}-repo"
            elif r["requer"] and not os.path.exists(os.path.join(base, r["requer"])):
                verdict, note = "Ausente", f"{r['requer']} não existe neste checkout"
            else:
                key = (base, r["comando"])
                if key not in cache:
                    cache[key] = run_producer(r["comando"], base)
                out, secs, rc = cache[key]
                verdict, note = judge(r["espera"], out, secs, rc)
        val = r["valor"]
        try:
            val = float(val) if val else None
        except ValueError:
            pass
        rows.append([r["id"], r["linha"], r["afirmacao"], val, r["unidade"], r["status"],
                     verdict, note, r["comando"], r["espera"], r["oraculo"], r["repositorio"],
                     r["citado_em"], r["medido_em"]])
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="sounio_ledger.xlsx")
    ap.add_argument("--uhs-repo")
    ap.add_argument("--gri-repo")
    ap.add_argument("--verify", action="store_true")
    ap.add_argument("--no-github", action="store_true")
    args = ap.parse_args()
    for k in ("uhs_repo", "gri_repo"):
        if getattr(args, k):
            setattr(args, k, os.path.abspath(getattr(args, k)))

    head = subprocess.run(["git", "-C", ROOT, "rev-parse", "--short", "HEAD"],
                          capture_output=True, text=True).stdout.strip()
    now = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    wb = Workbook()

    # Resumo first; filled after the other sheets exist
    ws0 = wb.active
    ws0.title = "Resumo"

    lines = research_rows(args)
    wsL = sheet(wb, "Linhas de pesquisa",
                ["ID", "Linha", "Repositório", "Caminho", "Caminho existe?", "Artefato principal",
                 "Estado medido", "Medido em", "Próximo passo", "Prazo", "Paper ou destino"],
                lines, [6, 40, 26, 34, 10, 36, 44, 11, 40, 12, 28])
    for row in wsL.iter_rows(min_row=2, max_row=wsL.max_row):
        for c in (row[8], row[9], row[10]):
            if c.value in (None, ""):
                c.fill = INPUT

    issue_note = ""
    if args.no_github:
        issues = []
        issue_note = "--no-github: aba vazia"
    else:
        try:
            issues = issue_rows()
        except Exception as e:  # report, never fabricate
            issues, issue_note = [], f"GitHub indisponível: {e}"
    wsI = sheet(wb, "Issues e PRs",
                ["#", "Tipo", "Título", "Área", "Labels", "Autor", "Aberto em", "Idade (dias)",
                 "Prioridade", "Link"], issues, [7, 7, 70, 13, 26, 16, 11, 9, 10, 46])
    for i in range(2, wsI.max_row + 1):
        wsI.cell(i, 7).number_format = "yyyy-mm-dd"
        wsI.cell(i, 8).value = f"=TODAY()-G{i}"
        wsI.cell(i, 9).fill = INPUT
    if issue_note:
        wsI["A2"] = issue_note

    lim = defect_rows() + kl_rows() + gap_rows(args)
    wsD = sheet(wb, "Limitações", ["ID", "Origem / área", "Descrição", "Motor", "Estado",
                                   "Rastreio", "Reprodução"], lim, [7, 26, 62, 12, 10, 34, 52])

    res = results_rows(args)
    wsR = sheet(wb, "Resultados",
                ["ID", "Linha", "Afirmação", "Valor", "Unidade", "Status", "Verificação",
                 "Nota", "Comando produtor", "Saída esperada", "Oráculo", "Repositório",
                 "Citado em", "Medido em"],
                res, [6, 6, 50, 14, 9, 11, 16, 26, 50, 30, 30, 22, 30, 11])
    n = wsR.max_row
    wsR.conditional_formatting.add(f"G2:G{n}", CellIsRule(operator="equal", formula=['"Verificado"'], fill=GREEN))
    for bad in ("Falhou", "Ausente"):
        wsR.conditional_formatting.add(f"G2:G{n}", CellIsRule(operator="equal", formula=[f'"{bad}"'], fill=RED))
    wsR.conditional_formatting.add(f"F2:F{n}", CellIsRule(operator="equal", formula=['"retirado"'], fill=GREY))
    wsR["G1"].comment = Comment("Só preenchida com --verify: o gerador roda o comando produtor e "
                                "procura a saída esperada. 'Ausente' = arquivo produtor não existe "
                                "neste checkout (por exemplo, PR ainda não mergeado).", "build_ledger.py")

    # ---- Resumo, all formulas over the sheets
    ws0.append(["Sounio — visão geral", None])
    ws0.append([f"Gerado em {now} a partir de main@{head} por tools/ledger/build_ledger.py"
                + (" com --verify" if args.verify else " (sem --verify)"), None])
    ws0.append([None, None])
    L, I, D, R = "'Linhas de pesquisa'", "'Issues e PRs'", "Limitações", "Resultados"
    rows = [
        ("Linhas de pesquisa", f"=COUNTA({L}!A2:A1000)"),
        ("  com caminho inexistente", f'=COUNTIF({L}!E2:E1000,"NÃO")'),
        ("  sem próximo passo definido", f'=COUNTBLANK({L}!I2:I{max(2, wsL.max_row)})'),
        ("Issues abertas", f'=COUNTIF({I}!B2:B1000,"Issue")'),
        ("PRs abertos", f'=COUNTIF({I}!B2:B1000,"PR")'),
        ("  idade mediana (dias)", f"=IFERROR(MEDIAN({I}!H2:H1000),0)"),
        ("Limitações e defeitos abertos", f'=COUNTIF({D}!E2:E1000,"aberto")'),
        ("Limitações e defeitos fechados", f'=COUNTIF({D}!E2:E1000,"fechado")'),
        ("Resultados registrados", f"=COUNTA({R}!A2:A1000)"),
        ("  válidos", f'=COUNTIF({R}!F2:F1000,"válido")'),
        ("  corrigidos", f'=COUNTIF({R}!F2:F1000,"corrigido")'),
        ("  retirados", f'=COUNTIF({R}!F2:F1000,"retirado")'),
        ("  verificados nesta geração", f'=COUNTIF({R}!G2:G1000,"Verificado")'),
        ("  falharam ou ausentes", f'=COUNTIF({R}!G2:G1000,"Falhou")+COUNTIF({R}!G2:G1000,"Ausente")'),
    ]
    for label, f in rows:
        ws0.append([label, f])
    ws0.append([None, None])
    ws0.append(["Células amarelas são para preencher à mão (próximo passo, prazo, paper, prioridade).", None])
    ws0.append(["Para atualizar: python3 tools/ledger/build_ledger.py --uhs-repo ../sounio-uhs-coupled --verify", None])
    ws0["A1"].font = Font(name=FONT, bold=True, size=14)
    for row in ws0.iter_rows(min_row=2):
        for c in row:
            if c.font != Font(name=FONT, bold=True, size=14):
                c.font = BODY
    ws0.column_dimensions["A"].width = 70
    ws0.column_dimensions["B"].width = 14

    wb.save(args.out)
    print(f"wrote {args.out}: {len(lines)} linhas, {len(issues)} issues/PRs, "
          f"{len(lim)} limitações, {len(res)} resultados"
          + (f" ({sum(1 for r in res if r[6] == 'Verificado')} verificados, "
             f"{sum(1 for r in res if r[6] in ('Falhou', 'Ausente'))} falharam/ausentes)" if args.verify else ""))
    if issue_note:
        print("note:", issue_note, file=sys.stderr)


if __name__ == "__main__":
    main()
